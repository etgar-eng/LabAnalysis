"""
main.py – PDF Lab Analysis Reader
Usage:
    python main.py --lab Bactochem --input bactochem/report.pdf --site "קידוח 5"
    python main.py --lab Aminolab  --input aminolab/report.pdf  --site "ביג פתח תקוה"
    python main.py --lab ALS       --input als/report.pdf       --multi true
"""

import argparse
import io
import logging
import os
import re
import sys
import warnings
from pathlib import Path

os.environ["TF_CPP_MIN_LOG_LEVEL"] = "3"
os.environ["TF_ENABLE_ONEDNN_OPTS"] = "0"
warnings.filterwarnings("ignore")

import camelot
import pdfplumber
import pandas as pd
from openpyxl import Workbook, load_workbook
from PIL import Image
from tensorflow.keras.models import load_model

from functions import Extractor, get_page_count, extract_sample_description
from identifier import predict_logo
from image_extractor import extract_images_hybrid_method, print_image_summary
from logging_config import setup_logging, suppress_warnings
from report import ReportBuilder, ReportMeta

pd.set_option("display.max_rows", 0)
pd.set_option("display.max_columns", 0)

logger = setup_logging(log_level=logging.INFO, log_file="lab_analysis.log")

# ── Table areas per lab ───────────────────────────────────────────────────────

# Aminolab VOC-only format (e.g. 11844.26): cover(1) + 3 VOC pages per check.
# None at index 0 is a placeholder so pages_per_check == 4 and offsets are correct.
_AMINOLAB_VOC_CONFIG: dict = {
    "tables_area": (None, None, None, None),
    "flavors":     ("stream", "stream", "stream", "stream"),
    "pages":       [None, "2", "3", "4"],
    "row_tol":     [7, 12],
}

_LAB_CONFIG: dict[str, dict] = {
    "ALS": {
        "tables_area": (
            ["30,50,410,680"], ["30,50,410,680"], None,
            ["30,50,410,680"], ["30,50,410,680"], ["30,50,410,680"], None,
        ),
        "pages":      ["3", "4", None, "6", "7", "8", "9"],
        "split_text": True,
        "strip_text": "\n",
    },
    "Bactochem": {
        "tables_area": (
            ["240,40,550,400"],
            ["240,30,550,750"],
            ["240,30,550,750"],
            ["240,30,550,750"],
            None,
        ),
        "pages":   ["1", "2", "3", "4", None],
        "row_tol": [7, 12],
    },
    "Aminolab": {
        # Standard mixed format (major elements on pg 2 + VOC on pgs 4-6).
        # pages_per_check = 8.  Page 6 switched from lattice → stream (no borders).
        "tables_area": (
            ["250,115,530,530"],  # pg 1 – cover (discarded later by content check)
            ["250,200,530,690"],  # pg 2 – ICP-MS major elements
            None,                  # pg 3 – notes, skip
            ["100,390,550,700"],  # pg 4 – VOC page 1
            ["30,110,460,680"],   # pg 5 – VOC page 2
            None,                  # pg 6 – VOC page 3 (full-page, stream)
            None,                  # pg 7 – notes, skip
            None,                  # pg 8 – disclaimer, skip
        ),
        "flavors": (
            "stream", "stream", None, "stream", "stream", "stream", None, None,
        ),
        "pages":   ["1", "2", None, "4", "5", "6", None, None],
        "row_tol": [7, 12],
    },
    "Element": {
        "tables_area": (
            ["30,100,600,600"],
            ["30,60,600,650"],
        ),
        "pages": ["2", "3"],
    },
}


# ── Helpers ───────────────────────────────────────────────────────────────────

def _print(msg: str):
    """Print safely, replacing characters that can't be encoded in the current terminal."""
    try:
        print(msg)
    except UnicodeEncodeError:
        print(msg.encode(sys.stdout.encoding or 'ascii', errors='replace').decode(sys.stdout.encoding or 'ascii', errors='replace'))


def _reset_excel(path: str):
    if os.path.exists(path):
        os.remove(path)
    wb = Workbook(); wb.save(path); wb.close()


def _remove_default_sheet(writer):
    if "Sheet" in writer.book.sheetnames and len(writer.book.sheetnames) > 1:
        writer.book.remove(writer.book["Sheet"])


def _last_real_page(pages: list):
    real = [p for p in pages if p is not None]
    return real[-1] if real else None


def _sheet_name(site: str, check_no: int, total_checks: int) -> str:
    """Return an Excel-safe tab name based on the borehole / sampling site.

    Rules enforced:
    - Strip characters forbidden in sheet names: \\ / * ? : [ ]
    - Truncate to 31 characters (Excel limit)
    - Append _N suffix when there are multiple checks in one PDF
    - Fall back to "Check{N}" when site is empty
    """
    base = re.sub(r'[\\/*?:\[\]]', '', site).strip()
    if not base:
        base = f"Check{check_no}"
    suffix = f"_{check_no}" if total_checks > 1 else ""
    # Leave room for the suffix
    max_base = 31 - len(suffix)
    return (base[:max_base] + suffix)


def _camelot_params(cfg: dict, page: str, area, flavor: str = "stream", is_first: bool = False) -> dict:
    params: dict = {"pages": page, "flavor": flavor}
    if area is not None:          # None → no area restriction (use full page)
        params["table_areas"] = area
    if flavor == "stream":
        if cfg.get("split_text") is not None:
            params["split_text"] = cfg["split_text"]
        if cfg.get("strip_text") is not None:
            params["strip_text"] = cfg["strip_text"]
        if cfg.get("row_tol"):
            params["row_tol"] = cfg["row_tol"][0] if is_first else cfg["row_tol"][-1]
    return params


# ── Multi-lab detector ────────────────────────────────────────────────────────

def detect_lab_pages(pdf_path: str, target_lab: str, model) -> list[str]:
    images = extract_images_hybrid_method(pdf_path)
    print_image_summary(images)
    class_names = ["ALS", "Bactochem", "Aminolab", "Element", "Yeda", "Yeda"]
    page_best: dict[int, tuple] = {}

    for i, img_info in enumerate(images):
        pil_img = Image.open(io.BytesIO(img_info.image_data))
        lab, conf = predict_logo(model, pil_img, class_names)
        page = img_info.page_number
        logger.info("Image %d (Page %d): %s (%.3f)", i, page, lab, conf)
        if page not in page_best or conf > page_best[page][1]:
            page_best[page] = (lab, conf)

    target_pages = sorted(p for p, (lab, _) in page_best.items() if lab == target_lab)

    if target_lab == "Bactochem" and target_pages:
        others = sorted(p for p, (lab, _) in page_best.items() if lab != target_lab)
        if others:
            target_pages = list(range(target_pages[0], others[0]))

    return [str(p) for p in target_pages]


# ── Extractor runner ──────────────────────────────────────────────────────────

class PdfLabAnalysisReader:
    """Read a lab PDF, extract tables with camelot, save to df_summary.xlsx."""

    def __init__(self, args):
        self.args = args
        self.lab  = args.lab
        self.pdf  = args.input
        lab_dir   = Path(os.getcwd()) / self.lab

        pdf_files = list(lab_dir.glob("*.pdf")) + list(lab_dir.glob("*.PDF"))
        if not pdf_files:
            raise FileNotFoundError(f"No PDF files found in {lab_dir}")

        self.page_count    = get_page_count(self.pdf)
        self.checks_number = self.page_count // 8 if self.lab == "Aminolab" else 1
        self.df_tables: list[pd.DataFrame] = []
        self.sample_dates: dict[str, str] = {}   # sheet_name → date extracted from PDF
        self.sheet_data:   dict[str, pd.DataFrame] = {}  # sheet_name → extracted DataFrame

        _reset_excel("df_summary.xlsx")
        _reset_excel(f"{self.lab}.xlsx")
        self._run()

    def _run(self):
        try:
            cfg  = _LAB_CONFIG[self.lab]

            # ── Aminolab: auto-detect VOC-only vs mixed format ─────────────
            # VOC-only reports (e.g. 11844.26) have 4 pages per check:
            #   cover(1) + 3 VOC pages.  The standard mixed format has 8 pages
            #   per check and puts major-element (ICP-MS) data on page 2.
            # Detection: if page 2 starts with a "VOC by Purge & Trap" section
            #   it is VOC-only; otherwise use the standard 8-page config.
            if self.lab == "Aminolab":
                with pdfplumber.open(self.pdf) as _pdf:
                    _p2 = (_pdf.pages[1].extract_text() or "") if len(_pdf.pages) >= 2 else ""
                if "VOC by Purge & Trap" in _p2:
                    logger.info("Aminolab: detected VOC-only format (4 pages/check)")
                    cfg = _AMINOLAB_VOC_CONFIG

            pages = cfg["pages"]

            if self.args.multi == "true":
                model  = load_model("lab_logo_classifier.h5")
                pages  = detect_lab_pages(self.pdf, self.lab, model)
                logger.info("Detected pages for %s: %s", self.lab, pages)

            areas_base   = list(cfg["tables_area"])
            flavors_base = list(cfg.get("flavors", ("stream",) * len(pages)))
            pages_per_check = len(pages)
            first_real   = next(p for p in pages if p is not None)
            last_real    = _last_real_page(pages)
            well_names = extract_sample_description(self.pdf)

            with pd.ExcelWriter("df_summary.xlsx", mode="a", engine="openpyxl",
                                if_sheet_exists="replace") as writer:

                for check_no, (_, (names, sample_date)) in enumerate(well_names.items(), start=1):
                    offset = (check_no - 1) * pages_per_check
                    last_page = str(int(last_real) + offset)

                    for page_base, area, flavor in zip(pages, areas_base, flavors_base):
                        if page_base is None:
                            continue
                        page = str(int(page_base) + offset)
                        is_first = page_base == first_real
                        params = _camelot_params(cfg, page, area, flavor or "stream", is_first)
                        try:
                            tables = camelot.read_pdf(self.pdf, **params)
                        except Exception as exc:
                            logger.warning("camelot error page %s: %s", page, exc)
                            continue

                        if tables:
                            self.df_tables.append(tables[0].df)

                        if page == last_page:
                            df_all = Extractor.extruct_col_from_lab(
                                self.df_tables, self.lab)
                            tab = names or _sheet_name(self.args.site, check_no, len(well_names))

                            # Keep a clean in-memory copy for report building
                            self.sheet_data[tab] = df_all.copy()

                            df_all.to_excel(writer, sheet_name=tab, index=False)
                            self.sample_dates[tab] = sample_date or ""
                            # Persist sample_date in E1 so df_summary is the
                            # single source of truth for downstream consumers
                            if sample_date:
                                writer.sheets[tab]["E1"] = sample_date
                            logger.info("Sheet written: '%s'", tab)
                            self.df_tables = []

                _remove_default_sheet(writer)

        except Exception as exc:
            logger.error("Extraction failed: %s", exc)
            import traceback; print(traceback.format_exc())


# ── CLI ───────────────────────────────────────────────────────────────────────

def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        description="PDF Lab Analysis → Water Authority Excel Report",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  python main.py --lab Bactochem --input bactochem/report.pdf
  python main.py --lab Aminolab  --input aminolab/report.pdf
  python main.py --lab ALS       --input als/report.pdf --multi true
        """,
    )
    p.add_argument("--lab",    choices=list(_LAB_CONFIG), required=True)
    p.add_argument("--input",  required=True, help="Path to PDF file")
    p.add_argument("--output", default=None,  help="Output name (no extension)")
    p.add_argument("--multi",  choices=["true", "false"], default="false")
    p.add_argument("--site",   default="",   help="Sampling site label")
    p.add_argument("--date",   default="",   help="Sampling date")
    p.add_argument("--time",   default="",   help="Sampling time")
    return p


def run(args):
    # 1 – load format + params
    extractor = Extractor(args)

    # 2 – extract raw tables from PDF
    logger.info("Extracting tables from %s", args.input)
    reader = PdfLabAnalysisReader(args)

    # 3 – build styled Excel report
    logger.info("Building Water Authority report")
    site        = args.site or Path(args.input).stem
    output_name = args.output or re.sub(r'[\\/*?:\[\]\s]', '_', site)
    report_dir  = Path(".")

    # Use in-memory sheet data (avoids stale df_summary.xlsx reads)
    sheet_names = list(reader.sheet_data.keys())

    for sheet_name in sheet_names:
        df_raw = reader.sheet_data[sheet_name]          # clean, freshly-extracted
        meta = ReportMeta(
            report_number=Path(args.input).stem,
            site=sheet_name if sheet_name else args.site,
            sampling_date=reader.sample_dates.get(sheet_name, "") or args.date,
            sampling_time=args.time,
            lab_name=args.lab,
        )
        builder = ReportBuilder(
            df_raw=df_raw,
            df_format=extractor.df_format,
            params=extractor.params,
            meta=meta,
        )
        safe = re.sub(r'[\\/*?:\[\]\s]', '_', sheet_name)
        out  = report_dir / f"{safe}_report.xlsx"
        builder.build(out)
        logger.info("Report written to %s", out)

    # 4 – plain mapped file (legacy compatibility), tabs named after site
    plain_path = Path(f"{output_name}.xlsx")
    _reset_excel(str(plain_path))   # ensure file exists before appending
    with suppress_warnings():
        with pd.ExcelWriter(str(plain_path), mode="a", engine="openpyxl",
                            if_sheet_exists="replace") as writer:
            for sheet_name in sheet_names:
                df_raw = reader.sheet_data[sheet_name]
                df_fmt = extractor.df_format.copy()
                df_fmt["values"] = df_fmt["test"].apply(
                    lambda t: extractor.search_for_value(df_raw, t)
                )

                df_fmt.to_excel(writer, index=False, sheet_name=sheet_name)

                date_val = reader.sample_dates.get(sheet_name, "")
                if date_val:
                    writer.sheets[sheet_name]["C3"] = date_val
            _remove_default_sheet(writer)

    _print(f"\n[OK] Raw data  -> df_summary.xlsx")
    _print(f"[OK] Mapped    -> {plain_path}")
    for sheet_name in sheet_names:
        safe = re.sub(r'[\\/*?:\[\]\s]', '_', sheet_name)
        _print(f"[OK] Report    -> {safe}_report.xlsx")


if __name__ == "__main__":
    args = build_parser().parse_args()
    try:
        run(args)
    except Exception as exc:
        print(f"Error: {exc}")
        import traceback; print(traceback.format_exc())
        sys.exit(1)
