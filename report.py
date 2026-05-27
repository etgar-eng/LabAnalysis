"""
report.py – Excel report generator for Israel Water Authority format.

Takes an extracted DataFrame (test / value / units columns) plus
lab metadata and writes a styled two-sheet Excel workbook:
  Sheet 1 – "דוח רשות המים"   full parameter table, colour-coded
  Sheet 2 – "סיכום תקינות"   compliance summary (exceedances + warnings)
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from regulations import Limit, compliance_status, get_limit

# ── Styling constants ────────────────────────────────────────────────────────

_THIN  = Side(style="thin",   color="BFBFBF")
_THIN_B = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_FILLS = {
    "title":   PatternFill("solid", fgColor="0A2342"),
    "header":  PatternFill("solid", fgColor="1F4E79"),
    "section": PatternFill("solid", fgColor="2E75B6"),
    "ok":      PatternFill("solid", fgColor="C6EFCE"),
    "warn":    PatternFill("solid", fgColor="FFD700"),
    "exceed":  PatternFill("solid", fgColor="FF6B6B"),
    "nd":      PatternFill("solid", fgColor="F2F2F2"),
    "white":   PatternFill("solid", fgColor="FFFFFF"),
    "alt1":    PatternFill("solid", fgColor="EBF5EB"),
    "alt2":    PatternFill("solid", fgColor="F0F8FF"),
    "alt3":    PatternFill("solid", fgColor="FDF5E6"),
}

_FONTS = {
    "title":  Font(name="Arial", bold=True, color="FFFFFF", size=14),
    "header": Font(name="Arial", bold=True, color="FFFFFF", size=10),
    "bold":   Font(name="Arial", bold=True, size=10),
    "body":   Font(name="Arial", size=10),
}

_STATUS_FILL = {
    "תקין":           _FILLS["ok"],
    "בגבול גילוי":   _FILLS["warn"],
    "לא ניתן לאמת":  _FILLS["warn"],
    "חריג!":          _FILLS["exceed"],
    "לא נבדק":        _FILLS["nd"],
}


# ── Metadata ─────────────────────────────────────────────────────────────────

@dataclass
class ReportMeta:
    report_number: str = ""
    lab_number: str = ""
    site: str = ""
    sampling_date: str = ""
    sampling_time: str = ""
    lab_name: str = ""
    sampler: str = ""
    sample_type: str = ""
    receipt_date: str = ""
    notes: str = ""


# ── Unit conversion ──────────────────────────────────────────────────────────

_CONVERSIONS: dict[tuple[str, str], float] = {
    ("mg/L",  "µg/L"): 1000.0,
    ("µg/L",  "mg/L"): 0.001,
    ("mg/l",  "mg/L"): 1.0,
    ("ug/l",  "µg/L"): 1.0,
    ("ug/L",  "µg/L"): 1.0,
    ("ppb",   "µg/L"): 1.0,
    ("ppm",   "mg/L"): 1.0,
}


def _coef(from_unit: str, to_unit: str) -> float:
    if from_unit == to_unit:
        return 1.0
    return _CONVERSIONS.get((from_unit, to_unit), 1.0)


def _extract_numeric(val) -> Optional[float]:
    if isinstance(val, (int, float)):
        return float(val)
    m = re.search(r"[\d,]+\.?\d*", str(val))
    if m:
        try:
            return float(m.group().replace(",", ""))
        except ValueError:
            pass
    return None


def _is_not_detected(val) -> bool:
    return isinstance(val, str) and "not detected" in val.lower()


# ── Writer helpers ───────────────────────────────────────────────────────────

def _cell(ws, row: int, col: int, value,
          fill=_FILLS["white"], font=_FONTS["body"],
          align: str = "center", wrap: bool = False):
    c = ws.cell(row=row, column=col, value=value)
    c.fill = fill
    c.font = font
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    c.border = _THIN_B
    return c


def _header_cell(ws, row: int, col: int, value, fill=_FILLS["header"]):
    return _cell(ws, row, col, value, fill=fill, font=_FONTS["header"],
                 align="center", wrap=True)


def _section_row(ws, row: int, title: str, ncols: int = 8):
    ws.merge_cells(f"A{row}:{get_column_letter(ncols)}{row}")
    c = ws.cell(row=row, column=1, value=title)
    c.fill = _FILLS["section"]
    c.font = _FONTS["header"]
    c.alignment = Alignment(horizontal="right", vertical="center")
    c.border = _THIN_B
    ws.row_dimensions[row].height = 22
    return row + 1


def _title_row(ws, row: int, text: str, ncols: int = 8):
    ws.merge_cells(f"A{row}:{get_column_letter(ncols)}{row}")
    c = ws.cell(row=row, column=1, value=text)
    c.fill = _FILLS["title"]
    c.font = _FONTS["title"]
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 32
    return row + 1


# ── Main report builder ──────────────────────────────────────────────────────

class ReportBuilder:
    """Build an Excel workbook from a normalised lab DataFrame.

    Parameters
    ----------
    df_raw : pd.DataFrame
        Normalised extraction result with columns ['test','value','units'].
        'test' contains letters-only text (Hebrew + ASCII) for synonym matching.
    df_format : pd.DataFrame
        format.csv table with columns ['test','units'].
    params : dict
        Lab-specific params.json mapping:
          {format_test_name: {"unit": target_unit, "values": [source_unit, synonym]}}
    meta : ReportMeta
        Report metadata (site, date, etc.)
    """

    def __init__(self, df_raw: pd.DataFrame, df_format: pd.DataFrame,
                 params: dict, meta: ReportMeta):
        self.df_raw = df_raw
        self.df_format = df_format
        self.params = params
        self.meta = meta
        self.wb = Workbook()

    # ── Public ───────────────────────────────────────────────────────────────

    def build(self, output_path: str | Path) -> Path:
        """Write the workbook to *output_path* and return the path."""
        self._build_report_sheet()
        self._build_summary_sheet()
        output_path = Path(output_path)
        self.wb.save(output_path)
        return output_path

    # ── Sheet 1 ──────────────────────────────────────────────────────────────

    def _build_report_sheet(self):
        ws = self.wb.active
        ws.title = "דוח רשות המים"
        ws.sheet_view.rightToLeft = True

        row = 1
        row = _title_row(ws, row,
                         f"דוח ניטור מי שתייה – פורמט רשות המים   |   {self.meta.site}")

        # Metadata block
        meta_pairs = [
            ("B", "אתר דיגום:", "C", self.meta.site),
            ("B", "מועד דיגום:", "C",
             f"{self.meta.sampling_date}  {self.meta.sampling_time}".strip()),
            ("B", "מעבדה:", "C", self.meta.lab_name),
            ("F", "מס' מעבדה:", "G", self.meta.lab_number),
            ("F", "דוגם:", "G", self.meta.sampler),
            ("F", "סוג דיגום:", "G", self.meta.sample_type),
            ("F", "תאריך קבלה:", "G", self.meta.receipt_date),
        ]
        for i, (bc, bl, vc, vl) in enumerate(meta_pairs):
            r = 2 + (i % 4)
            _cell(ws, r, _col(bc), bl, font=_FONTS["bold"],
                  align="right", fill=_FILLS["white"])
            _cell(ws, r, _col(vc), vl, font=_FONTS["body"],
                  align="right", fill=_FILLS["white"])
            ws.row_dimensions[r].height = 18

        # Column headers
        row = 7
        headers = [
            "פרמטר (רשות המים)", "פרמטר (מעבדה)",
            "תוצאה (מקורית)", "יחידה (מקורית)",
            "תוצאה (רש\"מ)", "יחידות (רש\"מ)",
            "גבול תקנות 2013", "סטטוס",
        ]
        for ci, h in enumerate(headers, 1):
            _header_cell(ws, row, ci, h)
        ws.row_dimensions[row].height = 36
        row += 1

        # Data rows grouped by section
        sections = self._group_by_section()
        alt_cycle = [_FILLS["alt1"], _FILLS["alt2"], _FILLS["alt3"]]

        for sec_idx, (section_title, tests) in enumerate(sections):
            row = _section_row(ws, row, f"► {section_title}")
            alt = alt_cycle[sec_idx % len(alt_cycle)]

            for i, test_name in enumerate(tests):
                row = self._write_data_row(ws, row, test_name,
                                           fill=alt if i % 2 == 0 else _FILLS["white"])

        # Column widths
        for ci, w in enumerate([32, 24, 14, 13, 12, 13, 16, 14], 1):
            ws.column_dimensions[get_column_letter(ci)].width = w
        ws.freeze_panes = "A8"

    def _write_data_row(self, ws, row: int, test_name: str, fill) -> int:
        param   = self.params.get(test_name, {})
        mapping = param.get("values", [])
        target_unit = param.get("unit", "—")

        lab_name_str = ""
        raw_val_str  = "—"
        raw_unit_str = "—"
        wra_val_str  = "—"
        wra_unit_str = target_unit
        limit_str    = str(get_limit(test_name))
        status       = "לא נבדק"

        if mapping and len(mapping) >= 2:
            source_unit = mapping[0]
            synonym     = mapping[1]
            lab_name_str = synonym

            found_val, found_unit = self._lookup(synonym, source_unit)

            if found_val is not None:
                raw_val_str  = str(found_val)
                raw_unit_str = found_unit

                if _is_not_detected(found_val):
                    wra_val_str = "Not Detected"
                    status      = compliance_status(test_name, "not detected")
                else:
                    n = _extract_numeric(found_val)
                    if n is not None:
                        c = _coef(source_unit, target_unit)
                        converted = n * c
                        wra_val_str = _fmt_number(converted)
                        status = compliance_status(test_name, converted)
                    else:
                        wra_val_str = str(found_val)
                        status = "תקין"

        sf = _STATUS_FILL.get(status, _FILLS["white"])
        row_data = [
            test_name, lab_name_str, raw_val_str, raw_unit_str,
            wra_val_str, wra_unit_str, limit_str, status,
        ]
        for ci, v in enumerate(row_data, 1):
            cell_fill = sf if ci == 8 else fill
            _cell(ws, row, ci, v, fill=cell_fill,
                  align="right" if ci <= 2 else "center")
        ws.row_dimensions[row].height = 18
        return row + 1

    # ── Sheet 2 ──────────────────────────────────────────────────────────────

    def _build_summary_sheet(self):
        ws = self.wb.create_sheet("סיכום תקינות")
        ws.sheet_view.rightToLeft = True

        row = 1
        _title_row(ws, row,
                   f"סיכום עמידה בתקנות בריאות העם (מי שתייה) 2013", ncols=5)

        # Headers
        row = 2
        for ci, h in enumerate(["פרמטר", "ערך מדוד", "יחידות",
                                  "גבול תקנות", "סטטוס / הערה"], 1):
            _header_cell(ws, row, ci, h)
        ws.row_dimensions[row].height = 28
        row += 1

        exceedances = []
        warnings    = []
        ok_notable  = []

        for test_name in self.df_format["test"]:
            param   = self.params.get(test_name, {})
            mapping = param.get("values", [])
            if not mapping or len(mapping) < 2:
                continue

            source_unit = mapping[0]
            synonym     = mapping[1]
            target_unit = param.get("unit", source_unit)

            found_val, _ = self._lookup(synonym, source_unit)
            if found_val is None:
                continue
            if _is_not_detected(found_val):
                continue

            n = _extract_numeric(found_val)
            if n is None:
                continue

            converted = n * _coef(source_unit, target_unit)
            status    = compliance_status(test_name, converted)
            lim       = get_limit(test_name)

            rec = (test_name, _fmt_number(converted), target_unit, str(lim), status)
            if status == "חריג!":
                exceedances.append(rec)
            elif status == "בגבול גילוי":
                warnings.append(rec)
            elif lim.value is not None:
                ok_notable.append(rec)

        # Write groups
        for group, label in [
            (exceedances, "🔴 חריגות מהתקן"),
            (warnings,    "🟡 ערכים בגבול / דורשים תשומת לב"),
            (ok_notable,  "✅ ערכים תקינים עם גבול מוגדר"),
        ]:
            if not group:
                continue
            ws.merge_cells(f"A{row}:E{row}")
            c = ws.cell(row=row, column=1, value=label)
            c.fill = _FILLS["section"]
            c.font = _FONTS["header"]
            c.alignment = Alignment(horizontal="right", vertical="center")
            c.border = _THIN_B
            ws.row_dimensions[row].height = 22
            row += 1

            for param, val, unit, lim_str, status in group:
                sf = _STATUS_FILL.get(status, _FILLS["white"])
                note = _compliance_note(param, val, unit, status)
                for ci, v in enumerate([param, val, unit, lim_str, note], 1):
                    f = sf if ci in (1, 5) else _FILLS["white"]
                    _cell(ws, row, ci, v, fill=f,
                          font=_FONTS["bold"] if ci == 1 else _FONTS["body"],
                          align="right" if ci in (1, 5) else "center",
                          wrap=(ci == 5))
                ws.row_dimensions[row].height = 40
                row += 1

        # Column widths
        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 14
        ws.column_dimensions["C"].width = 10
        ws.column_dimensions["D"].width = 18
        ws.column_dimensions["E"].width = 60

        # Legend
        row += 1
        ws.merge_cells(f"A{row}:E{row}")
        ws.cell(row=row, column=1, value="מקרא:").font = _FONTS["bold"]
        for i, (lf, lt) in enumerate([
            (_FILLS["exceed"], "חריגה מהתקן – נדרשת פעולה מיידית"),
            (_FILLS["warn"],   "בגבול / LOQ גבוה – נדרש מעקב"),
            (_FILLS["ok"],     "תקין"),
        ], row + 1):
            ws.merge_cells(f"A{i}:E{i}")
            c = ws.cell(row=i, column=1, value=lt)
            c.fill = lf
            c.font = _FONTS["body"]
            c.alignment = Alignment(horizontal="right")

        ws.freeze_panes = "A3"

    # ── Internals ─────────────────────────────────────────────────────────────

    def _lookup(self, synonym: str, source_unit: str):
        """Find synonym in df_raw; return (value, unit) or (None, None)."""
        if "test" not in self.df_raw.columns:
            return None, None

        mask = self.df_raw["test"].astype(str).str.contains(
            synonym, case=False, na=False)
        if not mask.any():
            return None, None

        row = self.df_raw[mask].iloc[0]
        val  = row.get("value")
        unit = row.get("units", source_unit)
        return val, unit

    def _group_by_section(self) -> list[tuple[str, list[str]]]:
        """Return (section_title, [test_names]) groups for Sheet 1."""
        sections = [
            ("כימיה כללית – אניונים וקטיונים", [
                "CALCIUM AS CA", "MAGNESIUM AS MG", "SODIUM AS NA",
                "POTASSIUM AS K", "CHLORIDE AS CL", "SULFATE AS SO4",
                "BICARBONATE AS HCO3", "NITRATE AS NO3", "FLUORIDE AS F",
                "BROMIDE AS BR", "PHOSPHATE AS PO4",
                "TOTAL ORGANIC CARBON (TOC)", "TOTAL NITROGEN AS N",
            ]),
            ("מתכות (ICP) – המרה לפי הצורך", [
                "SILVER (Ag)", "ALUMINUM (Al)", "ARSENIC AS AS",
                "BORON AS B", "BARIUM AS BA", "BERYLIUM (Be)",
                "ANTIMONY (Sb)", "CADMIUM AS CD", "COBALT AS CO",
                "CHROMIUM AS CR", "COPPER AS CU",
                "FERROUS IRON DISSOLVED AS FE", "MERCURY AS HG",
                "MANGANESE TOTAL AS MN", "NICKEL AS NI", "LEAD AS PB",
                "SELENIUM AS SE", "ZINC AS ZN",
            ]),
            ("תרכובות אורגניות נדיפות (VOC)", [
                "BENZENE", "ETHYL BENZENE", "TOLUENE", "XYLENE",
                "VINYL CHLORIDE (VC)", "CHLOROFORM",
                "TETRACHLOROETHYLENE (PCE)", "TRICHLOROETHYLENE (TCE)",
                "CARBON TETRACHLORIDE", "DICHLOROETHANE1,2",
                "DICHLOROETHYLENE 1,1 (1,1 DCE)",
                "CIS 1,2 DICHLOROETHYLENE (C1,2 DCE)",
                "TRANS-1,2 DICHLOROETHYLENE (T1,2 DCE)",
                "METHYLENE CHLORIDE (DCM)", "1,2-DICHLOROPROPANE",
                "1,1,2-TRICHLOROETHANE (1,1,2 TCA)",
                "TRICHLORO ETHANE 1,1,1 (1,1,1 TCA)",
                "ETHYLENE DIBROMIDE (EDB)", "MONOCHLOROBENZENE",
                "STYRENE", "BROMOBENZENE",
                "1,2,3-TRICHLOROPROPANE (TCP)",
                "1,3,5 TRIMETHYLBENZENE (MESITYLENE)",
                "1,2,4 TRIMETHYLBENZENE",
                "1,4 DICHLOROBENZENE (1,4 DCB)",
                "1,2 DICHLOROBENZENE (ODCB)",
                "1,2-DIBROMO-3-CHLOROPROPANE (DBCP)",
                "1,2,4-TRICHLOROBENZENE", "NAPTHALENE",
                "1,2,3-TRICHLOROBENZENE", "1,4 DIOXANE",
                "DIBROMOMETHANE", "DIBROMOCHLOROMETHANE",
                "TRIBROMOMETHANE (BROMOFORM)",
                "ACETONE", "BROMOMETHANE", "CHLOROMETHANE",
                "CHLOROETHANE", "DICHLORODIFLUOROMETHANE",
                "1,1-DICHLOROETHANE (1,1 DCA)", "1,1-DICHLOROPROPENE",
                "1,3-DICHLOROBENZENE", "MTBE",
            ]),
        ]
        # Only include tests that exist in df_format
        format_tests = set(self.df_format["test"])
        return [
            (title, [t for t in tests if t in format_tests])
            for title, tests in sections
        ]


# ── Helpers ──────────────────────────────────────────────────────────────────

def _col(letter: str) -> int:
    return ord(letter.upper()) - ord("A") + 1


def _fmt_number(n: float) -> str:
    if n == int(n) and abs(n) < 1e9:
        return str(int(n))
    return f"{n:.4g}"


def _compliance_note(test_name: str, val: str, unit: str, status: str) -> str:
    lim = get_limit(test_name)
    if status == "חריג!":
        return f"✗ חריגה! ערך {val} {unit} גבוה מהתקן ({lim})"
    if status == "בגבול גילוי":
        return f"⚠ ערך קרוב לגבול התקן ({lim}). יש לעקוב"
    return f"✓ תקין – גבול {lim}"
